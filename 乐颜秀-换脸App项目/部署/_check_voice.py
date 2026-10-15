# -*- coding: utf-8 -*-
from docx import Document
import re
import docx.oxml.ns as ns
from pathlib import Path

BASE = Path(r"REDACTED_LOCAL_USER_PATH\Documents\kimi\workspace\乐颜秀-换脸App项目\部署")
for name in ["算法安全自评估报告B（语音合成·备案提交版）.docx",
             "拟公示内容B（语音合成·备案提交版）.docx"]:
    d = Document(str(BASE / name))
    texts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                texts.append(c.text)
    full = "\n".join(texts)
    print("==", name)
    print("  图片数:", len(d.element.body.findall(".//" + ns.qn("a:blip"))),
          "| 待插入佐证:", full.count("待插入佐证"))
    for m in sorted(set(re.findall(r"【[^】]{1,24}】", full))):
        print("  占位:", m, full.count(m))
    for cap in [p.text for p in d.paragraphs if p.text.strip().startswith("图")]:
        print("  ", cap)
