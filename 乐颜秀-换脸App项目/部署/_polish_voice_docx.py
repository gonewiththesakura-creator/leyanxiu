# -*- coding: utf-8 -*-
"""完善语音备案材料B docx：
1) 4.5 投诉入口改为「意见反馈」（与 A 备案提交版截图口径一致）
2) 填报总览表补上「特别提示」行（md 有、docx 漏）
"""
import copy
from docx import Document

P = r"REDACTED_LOCAL_USER_PATH\Documents\kimi\workspace\乐颜秀-换脸App项目\部署\深度合成备案材料B（语音合成）.docx"
doc = Document(P)

# 1) 投诉入口口径统一
n = 0
for p in doc.paragraphs:
    if "「投诉举报」入口（我的 → 帮助与反馈）" in p.text:
        for r in p.runs:
            if "「投诉举报」入口（我的 → 帮助与反馈）" in r.text:
                r.text = r.text.replace(
                    "「投诉举报」入口（我的 → 帮助与反馈）",
                    "投诉举报入口（我的 → 意见反馈）")
                n += 1
print("投诉入口替换:", n)

# 2) 总览表补「特别提示」行
tb = doc.tables[0]
last = tb.rows[-1]
new_tr = copy.deepcopy(last._tr)
last._tr.addnext(new_tr)
from docx.table import _Row
row = _Row(new_tr, tb)
row.cells[0].paragraphs[0].runs[0].text = "特别提示"
row.cells[1].paragraphs[0].runs[0].text = (
    "声音复刻类功能为监管重点关注对象，本备案单独成册；"
    "即使被问询补正，也不影响 A 备案（图像视频）进度。"
)
# 清掉多余 run（深拷贝带来的原文本）
for c in row.cells:
    runs = c.paragraphs[0].runs
    for extra in runs[1:]:
        extra._element.getparent().remove(extra._element)
print("总览表现有行数:", len(tb.rows))

doc.save(P)
print("saved")
