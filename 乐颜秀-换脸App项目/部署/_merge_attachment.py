# -*- coding: utf-8 -*-
"""把落实算法安全主体责任基本情况合并到自评估报告备案版 docx 末尾（附件）。"""
import copy
from pathlib import Path
from docx import Document
from docx.enum.text import WD_BREAK

BASE = Path(__file__).parent
main = Document(BASE / "算法安全自评估报告（备案提交版）.docx")
att = Document(BASE / "落实算法安全主体责任基本情况（填报稿）.docx")

p = main.add_paragraph()
run = p.add_run("附件：")
run.add_break(WD_BREAK.PAGE)

body = main.element.body
sectPr = body.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr")
for el in att.element.body:
    if el.tag.split("}")[1] in ("p", "tbl"):
        sectPr.addprevious(copy.deepcopy(el))

out = BASE / "算法安全自评估报告（备案提交版-含主体责任附件）.docx"
main.save(out)
print("OK:", out.name)
