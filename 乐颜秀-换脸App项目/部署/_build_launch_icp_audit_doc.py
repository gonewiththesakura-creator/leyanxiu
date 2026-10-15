from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "乐颜秀App上线与ICP经营许可对抗式审核-执行版-2026-08-28.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
POSITIVE = "1F3A5F"
CAUTION = "7A5A00"
RISK = "9B1C1C"
WHITE = "FFFFFF"
BLACK = "111111"
MUTED = "5F6368"
BORDER = "CBD5E1"

FONT_ASCII = "Calibri"
FONT_CJK = "Microsoft YaHei"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def set_run_font(run, size=11, color=BLACK, bold=False, italic=False, ascii_font=FONT_ASCII, cjk_font=FONT_CJK):
    run.font.name = ascii_font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), cjk_font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_cell_margins(cell, margins=CELL_MARGINS_DXA):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in margins.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    assert sum(widths) == TABLE_WIDTH_DXA, (widths, sum(widths))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_paragraph_border_bottom(paragraph, color=BLUE, size="14", space="5"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])
    set_run_font(run, size=8.5, color=MUTED)


def add_num_pages_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " NUMPAGES "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])
    set_run_font(run, size=8.5, color=MUTED)


def set_cell_text(cell, text, bold=False, color=BLACK, size=9, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(str(text))
    set_run_font(run, size=size, color=color, bold=bold)
    return p


def add_table(doc, headers, rows, widths, font_size=8.6, status_col=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    keep_row_together(hdr)
    for idx, header in enumerate(headers):
        set_cell_shading(hdr.cells[idx], LIGHT_BLUE)
        set_cell_text(hdr.cells[idx], header, bold=True, color=INK, size=8.7)
    for row_data in rows:
        row = table.add_row()
        keep_row_together(row)
        for idx, value in enumerate(row_data):
            color = BLACK
            bold = False
            if status_col is not None and idx == status_col:
                bold = True
                value_str = str(value)
                if "未通过" in value_str or "阻断" in value_str or "No-Go" in value_str or "失败" in value_str:
                    color = RISK
                    set_cell_shading(row.cells[idx], "FCE8E6")
                elif "待确认" in value_str or "有条件" in value_str or "缺" in value_str:
                    color = CAUTION
                    set_cell_shading(row.cells[idx], "FFF4CE")
                elif "通过" in value_str or "完成" in value_str:
                    color = POSITIVE
                    set_cell_shading(row.cells[idx], "EAF0F7")
            set_cell_text(row.cells[idx], value, bold=bold, color=color, size=font_size)
        set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(doc, title, body, severity="risk"):
    fill = {"risk": "FCE8E6", "caution": "FFF4CE", "info": CALLOUT}[severity]
    title_color = {"risk": RISK, "caution": CAUTION, "info": POSITIVE}[severity]
    table = doc.add_table(rows=1, cols=1)
    set_repeat_table_header(table.rows[0])
    keep_row_together(table.rows[0])
    set_table_geometry(table, [TABLE_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11.5, color=title_color, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.2, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_para(doc, text="", size=11, color=BLACK, bold=False, italic=False, after=6, before=0, align=WD_ALIGN_PARAGRAPH.LEFT, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    p.paragraph_format.space_after = Pt(0)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT_ASCII)
    r_fonts.set(qn("w:hAnsi"), FONT_ASCII)
    r_fonts.set(qn("w:eastAsia"), FONT_CJK)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    r_pr.extend([r_fonts, color, underline, size])
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def define_numbering(doc):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    abstract_id = max(existing_abs or [0]) + 1
    num_id = max(existing_num or [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1.")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    lvl.extend([start, num_fmt, lvl_text, suff, p_pr])
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_numbered(doc, text, num_id):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    p_pr.append(num_pr)
    r = p.add_run(text)
    set_run_font(r, size=11, color=BLACK)
    return p


def set_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_ASCII
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_ASCII)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_ASCII)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading_tokens = {
        1: (16, BLUE, 18, 10),
        2: (13, BLUE, 14, 7),
        3: (12, DARK_BLUE, 10, 5),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = styles[f"Heading {level}"]
        style.font.name = FONT_ASCII
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_ASCII)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_ASCII)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True


def set_document_properties(doc):
    props = doc.core_properties
    props.title = "乐颜秀 App 上线与 ICP 经营许可对抗式审核（执行版）"
    props.subject = "App 上线、ICP 经营许可、生成式 AI、深度合成、人脸信息与安全合规"
    props.author = "乐颜秀项目组"
    props.keywords = "乐颜秀, App上线, ICP经营许可证, APP备案, AI合规, 深度合成, 人脸信息"
    props.comments = "基于截至 2026-08-28 的项目代码、公网测试版、既有材料及公开官方规则形成。"


def setup_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("乐颜秀｜上线与许可审查")
    set_run_font(r, size=8.5, color=MUTED, bold=True)
    r2 = hp.add_run("    执行版 2026-08-28")
    set_run_font(r2, size=8.5, color=MUTED)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_after = Pt(0)
    r3 = fp.add_run("内部执行文件  |  乐颜秀  |  证据基准 2026-08-28")
    set_run_font(r3, size=8.5, color=MUTED)
    return section


def build_document():
    doc = Document()
    set_document_properties(doc)
    set_styles(doc)
    setup_page(doc)
    num_id = define_numbering(doc)

    # First-page memo masthead.
    add_para(doc, "LAUNCH / LICENSE ADVERSARIAL REVIEW", size=9.5, color=BLUE, bold=True, after=5, before=8)
    add_para(doc, "乐颜秀 App 上线与 ICP 经营许可", size=23, color=BLACK, bold=True, after=3)
    title_rule = add_para(doc, "对抗式审核 · 执行版", size=14, color=MUTED, bold=False, after=14)
    set_paragraph_border_bottom(title_rule, color=BLUE, size="16", space="7")

    metadata = [
        ("审核对象", "REDACTED_COMPANY_NAME / 乐颜秀项目"),
        ("审核基准日", "2026 年 8 月 28 日"),
        ("证据范围", "当前代码、公网测试版、老板修改意见、既有交接与许可材料、现行公开官方规则"),
        ("审查方法", "先质疑需求 → 删除非必要上线范围 → 简化合规闭环 → 再做自动化和扩张"),
        ("当前决定", "NO-GO：可继续白名单/封闭测试，不得公众付费上线或提交正式商店审核"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.1
        lr = p.add_run(f"{label}：")
        set_run_font(lr, size=10.5, color=BLACK, bold=True)
        vr = p.add_run(value)
        set_run_font(vr, size=10.5, color=BLACK)

    add_callout(
        doc,
        "结论：当前不能正式上线",
        "阻断项不是 UI，而是人脸照片公开存储、无用户/任务归属鉴权、无可验证删除机制、无真实支付订单闭环、无正式 App 包与 APP 备案证据、无生产版隐私/单独同意流程、AI 标识与模型备案信息未形成完整闭环，同时生产依赖仍存在高危/严重漏洞。任何一项都足以单独阻断上线。",
        "risk",
    )

    add_heading(doc, "1. 一页决策", 1)
    add_table(
        doc,
        ["上线闸门", "当前状态", "决定依据", "放行条件"],
        [
            ("G1 主体与许可", "未通过", "注册资本表面满足省内增值电信业务最低 100 万元；但经营范围缺少电信业务字样，是否需 B25 尚未取得北京属地业务界定。", "完成业务界定；按结论取得所需备案/许可；主体、域名、服务器和商店账号一致。"),
            ("G2 隐私与 AI 合规", "未通过", "处理并经互联网传输人脸照片；未见首次启动同意、敏感个人信息单独同意、撤回/删除入口、模型备案信息公示。", "完成隐私政策、单独同意、PIPIA、第三方清单、撤回删除、AI 显/隐式标识和模型备案/登记证据。"),
            ("G3 生产安全与质量", "未通过", "OSS 公共读、任务无归属校验、CORS=*、无可验证 TTL 删除；依赖审计有严重/高危项；代码仓未形成干净可追溯发布基线。", "私有存储+签名 URL+鉴权+删除审计；漏洞清零或有批准例外；测试门槛达标；发布包可追溯。"),
            ("G4 商店、支付与运营", "未通过", "现状为 React H5 测试版，无 iOS/Android 原生包、签名、商店资料；会员和次数为本地演示逻辑，无订单、凭证、退款。", "完成原生包、APP 备案、正式支付、订单台账、退款客服、商店素材与审核账号。"),
        ],
        [1700, 1100, 3100, 3460],
        font_size=8.2,
        status_col=1,
    )

    add_heading(doc, "最小可上线产品（先删后加）", 2)
    add_table(
        doc,
        ["保留", "暂不上线", "原因"],
        [
            ("通过验收的核心视频模板；通过验收的少量图片能力；作品下载/删除；客服与投诉；合规分享文案。", "未通过样本验收的图片效果、照片跳舞扩展、玄学/生活玩法、伪榜单/伪热度、未授权音乐/人像/竞品素材、社区/用户发布。", "首版目标是“用户上传自己的照片，得到一个可用且可删除的 AI 作品”；功能数量、炫技动画和榜单都不是上线必要条件。"),
        ],
        [3000, 3300, 3060],
        font_size=8.7,
    )

    add_heading(doc, "2. P0 阻断项清单", 1)
    add_para(doc, "规则：P0 未全部关闭前，不开放公众注册、不接真实付费、不向应用商店提交正式审核。", size=10.5, color=RISK, bold=True, after=8)
    p0_rows = [
        ("P0-01", "关闭公开照片链路", "技术", "上传原图、生成结果、任务台账全部改为私有；仅返回短时签名 URL；历史公共对象盘点并迁移/删除。", "OSS ACL 报告、抽样 URL 匿名访问为 403、迁移清单"),
        ("P0-02", "用户与任务归属鉴权", "技术", "每个生成/查询/下载/删除请求绑定 userId 与 orderId；越权访问返回 403；禁止仅凭 taskId 查询。", "接口自动化越权测试 + 审计日志"),
        ("P0-03", "兑现删除承诺", "技术/法务", "把“24 小时清理”做成真实 TTL 作业；失败告警；用户可立即删除；原图、缓存、结果、第三方副本同步处理。", "定时任务日志、删除回执、抽样复测"),
        ("P0-04", "人脸信息单独同意", "法务/产品", "在上传前独立弹窗告知主体、目的、方式、保存期、第三方、影响与权利；明确本人/已获授权；可撤回。", "交互截图、同意记录字段、版本号、PIPIA"),
        ("P0-05", "AI 合成标识闭环", "技术/产品", "每张图片/每段视频增加用户可感知标识，并写入文件元数据的隐式标识；不得提供去标工具；分享导出仍保留。", "导出样本、元数据检测报告、UI 截图"),
        ("P0-06", "内容安全与权利库", "运营/法务", "上传前/生成后审核；拦截色情、暴力、未授权公众人物、未成年人高风险场景；建立音乐、模板、人像、字体授权台账。", "策略清单、命中样例、授权合同/采购凭证"),
        ("P0-07", "生产版协议与权利入口", "法务/产品", "隐私政策、用户协议、第三方共享清单、算法/模型公示、投诉举报、作品删除、账号注销、撤回同意均可从 App 内访问。", "正式 URL、页面截图、工单 SLA"),
        ("P0-08", "支付与订单", "产品/技术/财务", "删除“测试次数/补充次数/演示不扣款”；未接真支付前隐藏会员；上线付费后需订单、验签、幂等、恢复购买、退款和对账。", "沙盒+生产订单测试、退款用例、对账单"),
        ("P0-09", "正式包与备案", "公司/技术", "确定 iOS/Android/小程序首发渠道；完成包名、签名、隐私权限、APP/小程序备案；正式域名替换测试域名。", "备案号、安装包、签名指纹、商店后台截图"),
        ("P0-10", "依赖与发布基线", "技术", "修复生产依赖严重/高危漏洞；解决 lint 门禁；建立测试脚本；冻结干净 release tag；密钥轮换并清理 .env 备份。", "审计报告、CI 绿灯、tag/SBOM、密钥轮换记录"),
    ]
    add_table(doc, ["编号", "必须完成", "主责", "验收标准", "必须留下的证据"], p0_rows, [720, 1450, 800, 3440, 2950], font_size=7.65)

    add_heading(doc, "对现状的证据性判断", 2)
    evidence_rows = [
        ("公网形态", "已核验", "腾讯云 CloudBase 测试域名，进入前明确提示仅供开发测试；不能作为正式生产域名。"),
        ("真实模型", "部分已核验", "健康接口显示视频 provider=aliyun、图片 provider=viapi、文本 provider=deepseek；但“接口能返回”不等于每个效果质量合格或合规完成。"),
        ("视频模板", "已核验", "当前健康接口显示 4 个视频模板；首页仅首个视频处于播放状态。"),
        ("图片模板", "存在口径差异", "健康接口 photoTemplates=0，但前端展示多组图片效果；需按每个效果建立“前端入口→后端接口→供应商模型→验收报告”映射。"),
        ("构建与静态检查", "有条件", "前端 production build 通过；lint 有 8 个错误。后端语法检查通过。"),
        ("依赖风险", "未通过", "2026-08-28 审计快照：前端 2 个高危；后端 2 个严重、4 个高危、7 个中危。须升级或形成经批准的不可达/缓解证明。"),
        ("版本控制", "未通过", "main 分支存在大量修改、删除和未跟踪文件；当前 HEAD 不能代表公网生产版本。"),
    ]
    add_table(doc, ["对象", "状态", "审核结论"], evidence_rows, [1500, 1300, 6560], font_size=8.5, status_col=1)

    add_heading(doc, "3. ICP 经营许可证专项预审", 1)
    add_callout(
        doc,
        "先纠正概念：APP 备案 ≠ ICP 备案 ≠ ICP 经营许可证",
        "APP 备案是境内 App 提供互联网信息服务的基础手续；ICP备案对应非经营性互联网信息服务；俗称“ICP 证”的 B25 信息服务业务经营许可证，仅在实际业务落入增值电信业务且属于经营性时办理。付费 App 不是当然都要办 ICP 证，是否需要必须按真实业务形态界定。",
        "info",
    )

    icp_rows = [
        ("申请主体", "REDACTED_COMPANY_NAME", "有条件通过", "核对最新营业执照、法定代表人、股东穿透、外资成分和失信情况。"),
        ("注册资本", "100 万元", "表面通过", "达到省内增值电信业务最低限额，不能再下调；跨地区经营许可最低 1000 万元。"),
        ("经营范围", "有软件开发/信息技术咨询等，无“电信业务”字样", "未通过", "如申请许可证，先增加“许可项目：第二类增值电信业务”或主管部门认可的对应业务表述。"),
        ("业务类型", "自营 AI 图片/视频加工，计划收费；无第三方商户", "待属地确认", "最可能争议点是是否构成 B25 信息保护和处理服务；EDI 仅在做第三方交易平台时再评估。"),
        ("域名/备案", "CloudBase 测试域名", "未通过", "使用公司实名自有域名和境内接入，先完成网站/APP 备案，再按实际申请路径展示备案号。"),
        ("人员", "未见三岗位与社保证据", "材料缺失", "准备许可负责人、客服负责人、安全负责人及由申请主体缴纳的近期社保证明；以北京申请系统最新要求为准。"),
        ("业务与安全材料", "已有旧稿", "需重写", "旧稿中的模型、供应商、功能、域名、架构、存储方式与当前实现不一致，不得直接提交。"),
        ("前置审批", "营业执照含出版物零售；App 当前为 AI 影像", "需控范围", "首版不得混入新闻、出版、网络视听、教育、宗教等需专项审批服务；音乐/视频模板仍需版权授权。"),
    ]
    add_table(doc, ["审查项", "当前事实", "结论", "下一步"], icp_rows, [1650, 2900, 1250, 3560], font_size=8.05, status_col=2)

    add_heading(doc, "ICP 证决策路径", 2)
    add_numbered(doc, "形成一页真实业务说明：谁付费、买什么、是否有第三方发布/入驻、内容是否公开传播、数据如何处理、App/H5/小程序各自关系。", num_id)
    add_numbered(doc, "由公司主体向北京通信管理部门咨询业务界定并留存工单号/书面答复。不要让代理机构仅凭“收费”二字下结论。", num_id)
    add_numbered(doc, "若答复为无需许可证：保存界定证据，继续完成域名 ICP 备案、APP/小程序备案和其他 AI/数据合规。", num_id)
    add_numbered(doc, "若答复为需 B25：先变更经营范围，再以省内业务申请；准备 3 岗位人员、股权追溯、业务发展/技术方案、网络与信息安全材料。", num_id)
    add_numbered(doc, "若未来允许第三方作者/商家入驻、交易或用户公开发布内容，重新做 EDI、信息发布平台、内容审核和前置审批评估。", num_id)

    add_callout(doc, "ICP 专项最终判断", "资本不是当前障碍。真正障碍是业务界定未落定、经营范围不匹配、正式域名/备案缺失、人员与生产系统证据缺失。现阶段可以准备材料，但不应直接宣称“已经满足办证条件”。", "caution")

    add_heading(doc, "4. 公司方必须提供/办理", 1)
    company_rows = [
        ("C01", "确定首发渠道", "老板/产品负责人", "在 iOS、Android、微信小程序中选主路径；首发只保留一个收费闭环。", "书面决策 + 上线国家/地区"),
        ("C02", "工商与股权材料", "法人/行政", "最新营业执照、章程、法定代表人、所有股东逐级追溯、外资情况、公司印章。", "扫描件 + 股权树 + 真实性确认"),
        ("C03", "经营范围变更", "法人/行政", "仅在北京属地确认需申请 B25 后办理“第二类增值电信业务”等经营范围。", "新营业执照"),
        ("C04", "三岗位人员", "人事", "指定许可负责人、客服负责人、安全负责人；准备劳动关系和申请主体缴纳的近期社保证明。", "姓名/手机号/邮箱/社保证明"),
        ("C05", "官方数字资产", "法人/IT", "公司实名域名、境内云账号、ICP备案主体、Apple/Android/微信开发者主体统一。", "账号归属表 + 权限表"),
        ("C06", "收款与财务", "财务", "公司银行账户、商户号、税务开票、退款规则；iOS 数字权益使用 IAP。", "合同/商户号/对账责任人"),
        ("C07", "合同与供应商", "法务/采购", "与阿里云、腾讯云、DeepSeek/所用模型方确认企业主体、数据处理、存储地域、留存、训练使用、删除和 SLA。", "合同、DPA/条款快照、模型备案号"),
        ("C08", "内容授权", "运营/法务", "每个模板的视频、音乐、肖像、字体、封面均有可追溯商业授权；竞品截图仅作设计研究。", "资产编号→合同/授权链"),
        ("C09", "合规责任与客服", "老板/法务/客服", "确定个人信息负责人、投诉邮箱/电话、退款负责人、内容安全值班和事件上报链路。", "责任人名单 + SLA + 值班表"),
        ("C10", "预算批准", "老板/财务", "按法务、测试、云资源、模型调用、内容授权、商店账号、备案/许可服务分类批预算。", "预算上限 + 审批人"),
    ]
    add_table(doc, ["编号", "公司事项", "责任人", "必须提供/办理", "验收证据"], company_rows, [600, 1450, 1250, 3600, 2460], font_size=7.7)

    heading = add_heading(doc, "公司当前无需反复确认、执行方可直接推进", 2)
    heading.paragraph_format.page_break_before = True
    add_table(
        doc,
        ["执行方可自主完成", "只有触发下列情况才回报公司"],
        [
            ("代码安全整改、私有存储、签名 URL、鉴权、删除作业、日志、监控、自动化测试、漏洞升级、发布流水线、合规页面开发、资料截图整理、供应商能力映射。", "涉及真实付款、签署合同、购买资源、使用公章/法人账号、对外提交备案/许可、承诺保存期限、改变业务范围、引入新供应商或处理范围时。"),
        ],
        [4700, 4660],
        font_size=8.8,
    )

    add_heading(doc, "5. 执行方必须完成", 1)
    exec_rows = [
        ("E01", "生产架构", "将照片/视频/结果/任务从公共 URL 改为私有对象；签名 URL 最短有效；实现用户隔离、删除和备份清理。", "匿名访问 403；跨用户访问 403；删除后所有副本不可访问。"),
        ("E02", "身份与权益", "建立真实账号、服务端额度、订单、作品归属；不得以 localStorage 作为权益事实源。", "换设备仍一致；篡改本地存储不能增加权益。"),
        ("E03", "支付", "iOS IAP/Android 或小程序官方支付；服务端验签、幂等、补单、恢复购买、退款、对账。", "重复回调不重复发货；退款后权益回收；审计链完整。"),
        ("E04", "隐私同意", "首次启动基本同意；每次上传前敏感个人信息单独同意；第三方 AI 共享明确；撤回与删除可执行。", "协议版本、时间、目的、数据类型、供应商、用户动作均可追溯。"),
        ("E05", "AI 标识", "导出文件显式标识 + 元数据隐式标识；所有下载/分享通道保持。", "图片/视频样本经检测；二次导出仍保留。"),
        ("E06", "内容安全", "上传前/生成后检测；公众人物、未成年人、色情、暴恐、诈骗、侵权和去标识请求拦截；人工复核与申诉。", "至少 50 组对抗样本；命中、误杀、申诉记录可查。"),
        ("E07", "效果质量", "逐功能建立真实供应商、版本、提示词、成本、时延、成功率和样本集；未过线功能直接隐藏。", "核心功能至少 30 张差异化照片测试；成功率≥95%，严重变形/身份错配=0。"),
        ("E08", "发布工程", "清理未跟踪/敏感文件；release tag、SBOM、构建哈希、环境配置、回滚包、灰度开关。", "任何人可从 tag 复现同一安装包；生产密钥不进仓库。"),
        ("E09", "可观测性", "接口成功率、P95 时延、单次成本、队列积压、供应商错误、删除失败、内容安全命中告警。", "仪表盘 + 告警演练 + 日报样例。"),
        ("E10", "商店交付", "原生容器/客户端、包名、签名、权限说明、隐私标签、截图、审核账号、审核说明。", "真机测试报告；无占位页/演示字样；后端评审期间可用。"),
    ]
    add_table(doc, ["编号", "工作包", "必须做", "放行验收"], exec_rows, [600, 1300, 3820, 3640], font_size=7.8)

    add_heading(doc, "图片功能的第一性原理门槛", 2)
    add_para(doc, "图片功能不是“入口很多就算完成”。一个功能只有同时满足下面 5 项才允许出现在生产版。", size=10.5, bold=True, after=6)
    image_gate_rows = [
        ("1. 输入明确", "用户知道需要上传什么照片、预计得到什么结果；失败条件在上传前说清。"),
        ("2. 输出可识别", "结果确实呈现功能承诺，不靠标题或封面让用户脑补；与相邻功能的差异可肉眼辨认。"),
        ("3. 身份/画质安全", "人物身份不串脸、不变性别/年龄（除功能明确要求）、无明显肢体/五官错误、无未授权人像。"),
        ("4. 接口与成本可控", "真实接口、超时/重试/熔断、成本上限、供应商降级均可验证。"),
        ("5. 合规可删除", "单独同意、第三方告知、AI 标识、私有存储、作品删除、日志闭环全部存在。"),
    ]
    add_table(doc, ["门槛", "必须满足"], image_gate_rows, [1700, 7660], font_size=8.6)

    add_heading(doc, "6. 上线前测试与证据包", 1)
    test_rows = [
        ("功能", "所有公开入口均有真实可用后端；无空白页、占位图、假按钮；视频详情首屏展示模板并可播放。", "功能清单 + 逐项录屏"),
        ("兼容", "至少覆盖 2 个 iPhone、2 个主流 Android、弱网/断网/横竖屏/来电打断/后台恢复。", "机型矩阵 + 失败记录"),
        ("性能", "首页可交互、视频首帧、上传、生成排队、P50/P95 均有基线；超时有明确恢复。", "压测报告 + 仪表盘"),
        ("安全", "鉴权、越权、任意 URL、文件伪装、超大文件、恶意提示词、暴力重放、限流绕过、依赖漏洞。", "安全测试报告 + 修复复测"),
        ("隐私", "拒绝非必要权限仍可使用；撤回后停止处理；删除覆盖原图、结果、缓存、备份和供应商侧可控副本。", "数据流图 + 删除抽检"),
        ("支付", "购买、取消、失败、重复回调、补单、恢复购买、退款、换设备、对账全部覆盖。", "订单测试单 + 对账表"),
        ("内容", "高风险人像、色情、暴恐、未成年人、公众人物、侵权音乐、去 AI 标识请求均有策略。", "对抗样本结果 + 人工复核单"),
        ("运营", "客服、投诉、退款、下架模板、供应商故障、数据事件、舆情和监管问询有值班和时限。", "SOP + 演练记录"),
    ]
    add_table(doc, ["测试域", "最低放行标准", "证据"], test_rows, [1200, 5460, 2700], font_size=8.25)

    add_heading(doc, "当前自动化检查快照", 2)
    add_table(
        doc,
        ["检查", "结果", "结论"],
        [
            ("前端 production build", "通过", "可构建不代表可上线；只能证明 TypeScript/Vite 构建完成。"),
            ("前端 ESLint", "失败（8 项）", "纳入 CI 门禁后必须为 0 error。"),
            ("后端语法检查", "通过", "仍缺单元、集成、权限和删除测试。"),
            ("前端生产依赖审计", "2 个高危", "升级/替换或书面证明不可达并获安全负责人批准。"),
            ("后端生产依赖审计", "2 个严重、4 个高危、7 个中危", "生产发布前整改，重点清理旧 tcb-admin-node/request 依赖链并升级 Hono。"),
        ],
        [2400, 1900, 5060],
        font_size=8.6,
        status_col=1,
    )

    heading = add_heading(doc, "7. 30 天执行顺序", 1)
    heading.paragraph_format.page_break_before = True
    timeline_rows = [
        ("0-48 小时", "冻结发布范围；关闭/隔离公共上传；隐藏真实付费入口；删除假热度/未授权模板；指定公司 5 个责任人。", "老板 + 技术负责人", "范围冻结单、风险开关截图"),
        ("第 3-7 天", "私有存储、鉴权、删除、日志、密钥轮换；完成真实业务说明并向北京属地咨询 ICP 类别。", "技术 + 法人", "安全复测、咨询工单号"),
        ("第 1-2 周", "合规交互与协议；AI 标识；内容审核；核心功能样本验收；确定供应商与授权资产。", "产品/技术/法务/运营", "合规包、模型/资产台账、验收报告"),
        ("第 2-4 周", "正式域名和备案、原生包、支付沙盒/生产、商店材料、监控和应急演练。", "公司 + 执行方", "备案材料、安装包、支付与演练证据"),
        ("外部审批期", "按属地结论办理 APP 备案、ICP 证（如需）、生成式 AI 应用登记/深度合成算法备案（如适用）和商店审核。", "公司法定主体", "受理/批准编号；任何外部时限不作承诺"),
    ]
    add_table(doc, ["阶段", "必须完成", "主责", "退出条件"], timeline_rows, [1350, 4300, 1650, 2060], font_size=8.15)

    add_heading(doc, "法定/官方时限提示", 2)
    add_table(
        doc,
        ["事项", "官方口径", "项目排期含义"],
        [
            ("APP 备案", "材料齐全准确后，省级通信管理局应在 20 个工作日内予以备案。", "备案前先完成域名、主体、应用信息和接入商材料；不能把 20 个工作日当成从项目启动计算。"),
            ("增值电信业务许可", "法定审查上限为收到申请后 60 日；许可申请本身不收费。", "若属地界定需要 B25，需把经营范围变更和材料准备放在受理之前；代理服务费不是政府收费。"),
            ("商店审核", "平台不承诺固定通过时间；测试版不应直接提交正式商店。", "先用 TestFlight/封闭测试；必须提交最终功能、真实后台、审核账号和完整说明。"),
        ],
        [1900, 3300, 4160],
        font_size=8.35,
    )

    add_heading(doc, "预算只先建账，不先拍数字", 2)
    budget_rows = [
        ("监管与主体", "经营范围变更、CA/证照、必要的代理或律师复核", "公司", "政府许可申请不收费；第三方服务单独比价"),
        ("平台账号", "Apple/Android/微信企业开发者与签名证书", "公司", "按提交时官方价格"),
        ("云与模型", "对象存储、流量、数据库、日志、内容安全、AI 调用", "技术/财务", "设日/月成本上限与熔断"),
        ("版权", "视频、音乐、肖像、字体、封面商业授权", "运营/法务", "按资产编号逐项报价"),
        ("安全与合规", "渗透测试、隐私影响评估、必要的等保/测评、法律文本", "法务/安全", "先做范围界定再采购"),
        ("客服与运营", "退款、投诉、内容复核、数据事件值班", "客服/运营", "按 SLA 配置人力"),
    ]
    add_table(doc, ["预算桶", "包含内容", "负责人", "控制原则"], budget_rows, [1700, 3900, 1300, 2460], font_size=8.4)

    heading = add_heading(doc, "8. 最终放行签字表", 1)
    heading.paragraph_format.page_break_before = True
    gate_rows = [
        ("主体/许可", "业务界定、所需备案/许可、主体和域名一致", "□通过  □不通过", "公司负责人：________  日期：____"),
        ("隐私/人脸", "单独同意、PIPIA、第三方、撤回、删除", "□通过  □不通过", "法务/隐私负责人：____  日期：____"),
        ("AI/内容", "模型备案/登记、显隐式标识、审核、授权", "□通过  □不通过", "内容安全负责人：____  日期：____"),
        ("技术/安全", "鉴权、私有存储、删除、漏洞、监控、回滚", "□通过  □不通过", "技术负责人：________  日期：____"),
        ("质量/商店", "功能验收、真机、正式包、备案号、审核资料", "□通过  □不通过", "产品负责人：________  日期：____"),
        ("支付/客服", "订单、验签、退款、对账、投诉 SLA", "□通过  □不通过", "财务/客服：__________  日期：____"),
    ]
    add_table(doc, ["闸门", "必须证明", "结论", "签字"], gate_rows, [1500, 3900, 1500, 2460], font_size=8.4)
    add_callout(doc, "统一放行规则", "六个闸门全部签字通过，且 P0=0，才允许从封闭测试切换为公众上线。任何人不得以“先上再补”“供应商说可以”“竞品也这么做”替代证据。", "risk")

    add_heading(doc, "9. 官方依据与审核边界", 1)
    add_para(doc, "以下链接核对日期为 2026-08-28；主管机关、平台规则和供应商条款可能更新，正式提交前应再次复核。", size=9.5, color=MUTED, italic=True, after=6)
    sources = [
        ("工业和信息化部｜电信业务经营许可", "https://ythzxfw.miit.gov.cn/bssx/alx/dxhhlw/art/2025/art_5898effe0ba947d0bbfeff111b09f7fe.html"),
        ("通信管理局官方 FAQ｜无需许可情形与经营范围要求", "https://hubca.miit.gov.cn/zwgk/dxhhlwgl/art/2025/art_46268c23ed4f4a5191d9fc2f79dcf231.html"),
        ("工业和信息化部｜移动互联网应用程序备案通知", "https://www.miit.gov.cn/zwgk/zcwj/wjfb/tz/art/2023/art_920db564162e4312916a01bed6540ad8.html"),
        ("国家网信办｜互联网信息服务深度合成管理规定", "https://www.cac.gov.cn/2022-12/11/c_1672221949354811.htm"),
        ("国家网信办｜生成式人工智能服务管理暂行办法", "https://www.cac.gov.cn/2023-07/13/c_1690898327029107.htm"),
        ("国家网信办｜生成式人工智能服务已备案信息公告", "https://www.cac.gov.cn/2024-04/02/c_1713729983803145.htm"),
        ("国家网信办｜人工智能生成合成内容标识办法", "https://www.cac.gov.cn/2025-03/14/c_1743654684782215.htm"),
        ("国家网信办/公安部｜人脸识别技术应用安全管理办法", "https://www.cac.gov.cn/2025-03/21/c_1744174262156096.htm"),
        ("工业和信息化部｜个人信息保护法", "https://www.miit.gov.cn/jgsj/zfs/fl/art/2022/art_515a4b20c12f430eab54bb4f56d89f56.html"),
        ("Apple Developer｜App Review Guidelines", "https://developer.apple.com/app-store/review/guidelines/"),
    ]
    for idx, (label, url) in enumerate(sources, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.1
        r = p.add_run(f"{idx}. {label}：")
        set_run_font(r, size=9, color=BLACK, bold=True)
        add_hyperlink(p, "打开官方原文", url)

    add_heading(doc, "审核边界", 2)
    add_para(doc, "本文件是项目上线的操作性预审，不是律师正式法律意见，也不替代北京通信管理局、属地网信部门、应用商店或其他主管机关的最终认定。对 ICP 经营许可证，本文给出的是“先做属地业务界定、再决定是否申请”的控制结论；对生成式 AI、深度合成和人脸信息，必须以项目实际调用模型、功能、用户规模、数据流和传播能力逐项确认。", size=9.6, color=MUTED, after=8)

    add_para(doc, "文档版本：v1.0｜证据基准：2026-08-28｜建议复审节点：每次供应商、数据流、付费方式、上线渠道或公开传播能力发生变化时。", size=8.8, color=MUTED, italic=True, after=0)

    # Enable automatic field updates when opened in Word.
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_document())
